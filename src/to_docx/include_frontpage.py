from docx import Document
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
from src.constants import TEXT_TO_REPLACE_IN_FRONTPAGE
from bs4 import BeautifulSoup, Tag, NavigableString
from src.to_docx.preprocessing import remove_diff_marks

def scrap_replacements_from_html(front_page_html_file: str) -> dict:
    replacements = {}
    with open(front_page_html_file, 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'html.parser')
        soup = remove_diff_marks(soup)
        for key in TEXT_TO_REPLACE_IN_FRONTPAGE:
            element = soup.find(attrs={"data-replace": key})
            if element:
                if key == "DATE":
                    # Special handling for DATE to format it as needed
                    date_text = element.get_text(strip=True)
                    split_date = date_text.split('-')
                    YEAR = split_date[0] if len(split_date) > 0 else ''
                    MONTH = split_date[1] if len(split_date) > 1 else ''
                    replacements["yyyy"] = YEAR
                    replacements["-mm"] = f"-{MONTH}"
                else:
                    replacements[key] = element.get_text(strip=True)
    return replacements

def include_frontpage(output_path):
    def merge_docx(output_path: str, *input_paths: str):
        master = Document(input_paths[0])
        composer = Composer(master)
        for path in input_paths[1:]:
            p = master.add_paragraph()
            run = p.add_run()
            run.add_break(WD_BREAK.PAGE)
            
            composer.append(Document(path))
        composer.save(output_path)


    def replace_placeholders(file_path: str, replacements: dict, output_path: str = None):
        doc = Document(file_path)
        
        for paragraph in doc.paragraphs:
            for key in TEXT_TO_REPLACE_IN_FRONTPAGE:
                value = replacements.get(key, "")
                if key in paragraph.text:
                    for run in paragraph.runs:
                        run.text = run.text.replace(key, value)
        
        doc.save(output_path or file_path)

    def set_update_fields_on_open(doc_path: str):
        """Force to update fields when opening the document in Word."""
        doc = Document(doc_path)
        
        settings = doc.settings.element
        update_fields = parse_xml(
            f'<w:updateFields {nsdecls("w")} w:val="true"/>'
        )
        settings.append(update_fields)
        
        doc.save(doc_path)

    front_page_html = os.path.join(os.path.dirname(output_path), "html", "front-page.html")
    replacements = scrap_replacements_from_html(front_page_html)
    document = os.path.join(output_path, "document.docx")
    front_page_filled = os.path.join(output_path, "front_page_filled.docx")
    replace_placeholders("front_page_template.docx", replacements, front_page_filled)
    merge_docx(document, front_page_filled, document)
    # set_update_fields_on_open(document)
    os.remove(front_page_filled)
