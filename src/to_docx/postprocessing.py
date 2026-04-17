import subprocess
from cProfile import run
import copy
import re, os
import shutil
import tempfile
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.document import Document as Doc
from docx.text.paragraph import Paragraph, Run
from docx.text.run import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from src.constants import (
    BOLD_TAGS,
    EXAMPLE_NOTE_CLASSES,
    STYLES_FOR_EXAMPLES_NOTES,
    HANDLE_UNDERSCORE_CLASSES,
    HTML_BASIC_FORMAT_TAGS,
    ITALIC_TAGS,
    NEWLINE_PLACEHOLDER,
    NO_SPACE,
    PERM_CSS_FILE,
    STRIKETHROUGH_TAGS,
    SUBSCRIPT_TAGS,
    SUPERSCRIPT_TAGS,
    TAB_PLACEHOLDER,
    UNDERLINE_TAGS,
    WITH_SPACE,
)
from src.to_docx.include_frontpage import include_frontpage
from src.utils import get_bold_italic_underline_css_classes


# region Helpers
def iter_paragraphs(doc: Doc):
    """Yield all paragraphs in the document, including in tables, headers, footers."""
    for el in doc.element.iter():
        if el.tag == qn("w:p"):
            yield Paragraph(el, doc)


def format_references(doc: Doc):
    """Add a tab after the tag to ensure body text is correctly indented according to the `EX` style"""

    def insert_tab_after_reference_tag(paragraph, label: str):
        """Modify the paragraph's XML to add a tab after the label"""
        xml = paragraph._element
        total = ""
        found = False

        # Collect all <w:r> (run) elements regardless of hyperlink wrapping
        runs = xml.findall(".//" + qn("w:r"))

        for run in runs:
            label_component = run.find(qn("w:t"))

            if label_component is None:
                continue  # This isn't a reference if it doesn't have the label

            text = label_component.text or ""
            total += text

            if not found and label in total:
                # Label found; split text
                label_position = total.index(label)
                end_of_label_position = label_position + len(label)

                # Trim the run text to end at label
                num_preceding_chars = len(total) - len(text)
                num_current_run_chars = end_of_label_position - num_preceding_chars

                # Separate the reference's text between the label and the body
                tag_text = text[:num_current_run_chars]
                text_after_tag = text[num_current_run_chars:]

                # Set the label as the active text node
                label_component.text = tag_text

                # Get the hyperlink from the run
                parent_element = run.getparent()
                hyperlink = None
                if parent_element.tag == qn("w:hyperlink"):
                    hyperlink = parent_element

                # Get the insertion point for the tab (after the hyperlink)
                tab_insertion_point = run

                if hyperlink is not None:
                    tab_insertion_point = hyperlink

                # Create a new run to contain the tab outside the hyperlink
                tab_run = OxmlElement("w:r")
                tab = OxmlElement("w:tab")
                tab_run.append(tab)

                # Insert the tab run after the hyperlink
                xml.insert(xml.index(tab_insertion_point) + 1, tab_run)

                # Re-add text after label to the run
                if text_after_tag:
                    body_run = OxmlElement("w:r")
                    body_text = OxmlElement("w:t")
                    body_text.text = text_after_tag
                    body_run.append(body_text)
                    xml.insert(xml.index(tab_run) + 1, body_run)

                found = True
                break

    reference_regex = re.compile(r"^(\[(?:i\.)?\d{1,2}\])")

    for paragraph in doc.paragraphs:
        label = reference_regex.match(paragraph.text)
        if paragraph.style.name.startswith("EX") and label:
            label = label.group(0)
            insert_tab_after_reference_tag(paragraph, label)

    return doc


def format_tables(doc: Doc):
    """
    Ensures tables are formatted correctly by doing the following:
    - Set border thickness to 1/2 pt
    - Apply table header styles within table headers. Pandoc should be able to do this just like it can with body cells, but unfortunately, that isn't the case.
    """

    def set_borders(table: CT_Tbl):
        """Ensure table borders are set to 1/2 pt thickness"""
        # Remove existing borders
        existing_borders = table.tblPr.find(qn("w:tblBorders"))

        if existing_borders is not None:
            table.tblPr.remove(existing_borders)

        # Add new borders
        borders = OxmlElement("w:tblBorders")

        for border_location in [
            "top",
            "left",
            "bottom",
            "right",
            "insideH",
            "insideV",
        ]:
            border = OxmlElement(f"w:{border_location}")

            border.set(qn("w:val"), "single")  # Style: Single line
            border.set(
                qn("w:sz"), "4"
            )  # Size: 1/2 pt - size is measured in 1/8 pt increments, so 1/8 * 4 = 1/2 pt

            borders.append(border)

        table.tblPr.append(borders)

        return table

    def set_table_header_styles(table: CT_Tbl):
        """Apply the tagged styles within header cells to the contents of that cell, then delete the tag."""

        def get_stylename(first_para: BaseOxmlElement):
            """Returns the stylename defined in the first paragraph if applicable, otherwise, returns `None`"""
            text_elements = first_para.xpath(".//w:t")

            if not text_elements:
                return None

            tag_text: str = text_elements[0].text or ""

            TAG_REGEX = r"\{{3}([^\{\}]+)\}{3}"
            tag_match = re.match(TAG_REGEX, tag_text.strip())

            if not tag_match:
                return None  # No tag in this cell

            stylename: str = tag_match.group(1)

            return stylename

        def apply_style_to_para(para: BaseOxmlElement, stylename: str):
            """Applys the style provided to `stylename` to the paragraph `para`"""
            para_props = para.find(qn("w:pPr"))

            if para_props is None:
                para_props = OxmlElement("w:pPr")
                para.insert(0, para_props)

            for old_p_style in para_props.findall(qn("w:pStyle")):
                # Clear the existing pStyles to preserve the correct XML structure, thereby avoiding unreadable content errors when opening the Docx
                para_props.remove(old_p_style)

            para_style = OxmlElement("w:pStyle")
            para_style.set(qn("w:val"), stylename)

            para_props.insert(0, para_style)

            return para

        cells: list[BaseOxmlElement] = table.xpath(".//w:tc")

        for cell in cells:
            paras: list[BaseOxmlElement] = cell.xpath(".//w:p")

            if not paras:
                continue  # No paragraphs in this cell, nothing to do

            first_para = paras[0]
            stylename = get_stylename(first_para)

            if stylename is None:
                continue  # No tag in this cell

            is_remove_tag = len(paras) > 1

            if is_remove_tag:
                first_para.getparent().remove(first_para)
            else:
                # Maintain correct XML structure by clearing the tag, but keeping the paragraph
                for text in first_para.xpath(".//w:t"):
                    text.text = ""

            for para in paras[1:]:
                para = apply_style_to_para(para, stylename)

        return table

    for table in doc.tables:
        table = table._element

        table = set_borders(table)
        table = set_table_header_styles(table)

    return doc


def format_examples_and_notes(doc: Doc):
    """Add tabs in to examples and notes to ensure their bodies are indented properly"""

    def handle_paragraphs(paragraphs: list[Paragraph]):
        """Adds tabs after the example/note labels of all the examples and notes in the provided paragraphs"""
        ex_no_tan_regex = re.compile(r"^(EXAMPLE(?:\s\d{1,2})?:|NOTE(?:\s\d{1,2})?:)")

        def has_exno_markers(paragraph: Paragraph, index: int):
            """Return True when the paragraph contains invisible EX/NOTE boundary markers."""
            for run in paragraph.runs:
                if run.style and run.style.name == "InvisibleText":
                    if "startEXNO" in run.text or "endEXNO" in run.text:
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.paragraph_format.space_before = Pt(0)

                        prec_paragraph = paragraphs[index - 1] if index > 0 else None
                        if prec_paragraph:
                            prec_paragraph.paragraph_format.space_after = Pt(0)
                        return True
            return False

        in_example_or_note = False

        for paragraph, index in zip(paragraphs, range(len(paragraphs))):
            tag = ex_no_tan_regex.match(paragraph.text)
            contains_exno_markers = has_exno_markers(paragraph, index)

            if paragraph.style.name in EXAMPLE_NOTE_CLASSES and tag:
                # Start of example or note
                if not contains_exno_markers:
                    tag = tag.group(0)
                    if "~~~" in paragraph.text: 
                        tilde_text = "{{{InvisibleText}}}~~~{{{/InvisibleText}}}\t"

                        split_text = paragraph.text.split("~~~", 1)
                        body_text = f"{split_text[1]}" if len(split_text) > 1 else ""

                        paragraph.text = tag + tilde_text + body_text
                    else:
                        body_text = paragraph.text[len(tag) :]
                        paragraph.text = f"{tag}\t{body_text}"
                in_example_or_note = True
                continue  # Avoid indenting the reference label

            if in_example_or_note:
                # Continue example or note
                if paragraph.style.name in EXAMPLE_NOTE_CLASSES:
                    if not contains_exno_markers:
                        paragraph.text = f"\t{paragraph.text}"
                elif paragraph.style.name in STYLES_FOR_EXAMPLES_NOTES:
                    if paragraph.style.name == "Source Code":
                        # paragraph.style = "EX Source Code"
                        paragraph.paragraph_format.left_indent = Pt(
                            85.04
                        )  # 3 cm in points

                    if paragraph.style.name == "Compact":
                        # paragraph.style = "EX Compact"
                        paragraph.paragraph_format.left_indent = Pt(
                            103.68
                        )  # 3.66 cm in points
                else:
                    # No longer in example or note
                    in_example_or_note = False

            if (
                paragraph.style.name == "Source Code"
                or paragraph.style.name == "EX Source Code"
            ):
                # apply to all its runs the HTML-Sample style with no space suffix
                for run in paragraph.runs:
                    if run.style != "InvisibleText":
                        run.style = "HTML-Sample"

    document_paragraphs = list(iter_paragraphs(doc))

    handle_paragraphs(document_paragraphs)

    return doc


def merge_example_with_next_sourcecode(doc: Doc):
    """
    Merge EXAMPLE/NOTE paragraphs with the following Source Code paragraph,
    ONLY when the example paragraph contains just the label (e.g. "EXAMPLE 1:\t").
    This replicates what Word does when you press backspace between the two paragraphs.
    """
    # Match only paragraphs that contain just the label and optional whitespace
    ex_label_only_regex = re.compile(
        r"^(EXAMPLE(?:\s\d{1,2})?:|NOTE(?:\s\d{1,2})?:)(?:\s*~~~)?\s*$"
    )

    paragraphs = list(iter_paragraphs(doc))
    to_remove = []

    for i in range(len(paragraphs) - 1):
        current = paragraphs[i]
        next_para = paragraphs[i + 1]
        # Check if current is an example/note label-only paragraph
        if current.style.name in EXAMPLE_NOTE_CLASSES and ex_label_only_regex.match(
            current.text.strip()
        ):
            # Check if next paragraph is a Source Code block
            if next_para.style.name in ["Source Code", "EX Source Code"]:
                # Move all runs from the Source Code paragraph into the current paragraph
                for r_elem in list(next_para._p.findall(qn("w:r"))):
                    current._p.append(r_elem)

                # Mark the now-empty Source Code paragraph for removal
                to_remove.append(next_para._p)

    # Remove all marked paragraphs after the loop to avoid modifying the list while iterating
    for p_elem in to_remove:
        p_elem.getparent().remove(p_elem)

    return doc

def remove_startEXNO_invisible_text_when_in_the_same_paragraph(doc: Doc):
    """
        Remove the startEXNO invisible text when it is in the same paragraph as the endEXNO invisible text.
        When this happens this is a reference, which is very well formatted.
    """
    paragraphs = doc.paragraphs
    for paragraph in paragraphs:
        index_startEXNO = -1
        index_endEXNO = -1

        for run, index in zip(paragraph.runs, range(len(paragraph.runs))):
            if run.style and run.style.name == "InvisibleText":
                if "startEXNO" in run.text:
                    index_startEXNO = index
                if "endEXNO" in run.text:
                    index_endEXNO = index
        
        if index_startEXNO != -1 and index_endEXNO != -1:
            # remove these runs from the paragraph
            paragraph._element.remove(paragraph.runs[index_endEXNO]._r)
            paragraph._element.remove(paragraph.runs[index_startEXNO]._r)
    return doc


def replace_whitespace_placeholders_with_chars(doc: Doc):
    """Replace any tab and newline placeholders added during preprocessing with real tabs and newlines"""
    # Alternative style based on `Heading 1` for top-level annex headers. Removes the tab indent, since the heading body is on a newline after the heading label.
    TOP_LEVEL_HEADING_STYLE_FOR_NEWLINES = "Heading 8"

    def replace_tabs(string: str):
        """Replaces any tab placeholders in the provided string with tab characters."""
        return string.replace(TAB_PLACEHOLDER, "\t")

    for paragraph in iter_paragraphs(doc):
        if TAB_PLACEHOLDER in paragraph.text or NEWLINE_PLACEHOLDER in paragraph.text:
            for run in paragraph.runs:
                if NEWLINE_PLACEHOLDER not in run.text:
                    run.text = replace_tabs(run.text)

                else:
                    text = run.text
                    run.text = ""

                    text_parts = text.split(NEWLINE_PLACEHOLDER)

                    for i, part in enumerate(text_parts):
                        part = replace_tabs(part)  # Replace any tabs there may be
                        run.add_text(part)

                        if i < len(text_parts) - 1:
                            run.add_break(WD_BREAK.LINE)

                    paragraph.style = TOP_LEVEL_HEADING_STYLE_FOR_NEWLINES

    return doc


def add_tagged_styles_and_formatting(doc: Doc):
    """Re-add the styles and other formatting that were lost during conversion back into the text"""
    CSS_BOLD, CSS_ITALIC, _ = get_bold_italic_underline_css_classes(PERM_CSS_FILE)

    def handle_basic_formatting(run: Run, style_name: str):
        """Handles the formatting of sections of text tagged with basic HTML font formatting tags, like bold, italics, and so on."""
        if style_name in BOLD_TAGS:
            run.bold = True

        if style_name in ITALIC_TAGS:
            run.italic = True

        if style_name in UNDERLINE_TAGS:
            run.underline = True

        if style_name in STRIKETHROUGH_TAGS:
            run.font.strike = True

        if style_name in SUBSCRIPT_TAGS:
            run.font.subscript = True

        if style_name in SUPERSCRIPT_TAGS:
            run.font.superscript = True

        return run

    def add_hyperlink(paragraph: Paragraph, url: str, text: str):
        """Add a hyperlink to text within a paragraph"""
        hyperlink = OxmlElement("w:hyperlink")

        # Internal anchors are represented as LINK|#anchor and restored as w:anchor.
        if url.startswith("#") and len(url) > 1:
            hyperlink.set(qn("w:anchor"), url[1:])
        else:
            part = paragraph.part
            run_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
            hyperlink.set(qn("r:id"), run_id)

        # Create a new w:r run and add the Hyperlink style
        new_run = OxmlElement("w:r")

        run_properties = OxmlElement("w:rPr")
        run_style = OxmlElement("w:rStyle")

        run_style.set(qn("w:val"), "Hyperlink")
        run_properties.append(run_style)

        # Create a new w:t text
        w_text = OxmlElement("w:t")
        w_text.text = text

        # Add the text and the properties to the run
        new_run.append(run_properties)
        new_run.append(w_text)

        # Add the run to the hyperlink and the hyperlink to the paragraph
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)

        # Wrap XML run so callers can further format it (bold/italic/etc.)
        return Run(new_run, paragraph)

    def handle_code_block_tags_with_suffixes(para: Paragraph, stylename: str):
        """Strip the suffix from the `stylename`, then check the suffix to see whether the paragraph should have any space after it. By default the paragraph has space after it, so remove the space after the paragraphs for those that were tagged with the suffix indicating that no space should be added afterward."""
        if not (
            stylename == f"HTML_Sample/{NO_SPACE}"
            or stylename == f"HTML_Sample/{WITH_SPACE}"
        ):
            return (
                para,
                stylename,
            )  # Code block tags are in the format "HTML_Sample/{suffix}"

        stylename, suffix = stylename.split("/")

        # Convert internal tag name to actual Word style name
        stylename = "HTML-Sample"

        if suffix == NO_SPACE:
            para.paragraph_format.space_after = Pt(0)

        return para, stylename

    def apply_formatting_and_styling(run: Run, style_name: str):
        if run.style and run.style.name == "InvisibleText":
            return run  # Don't apply any formatting or styling to invisible text

        para: Paragraph = run._parent
        txt = run.text

        para, style_name = handle_code_block_tags_with_suffixes(para, style_name)

        if style_name in HTML_BASIC_FORMAT_TAGS:
            # Apply basic formatting, ex., bold, italics, etc
            run = handle_basic_formatting(run, style_name)

        elif style_name.startswith("LINK|") and txt.strip():
            # Replace the run with a hyperlink
            link = style_name.split("|", 1)[1]

            run.text = ""
            add_hyperlink(para, link, txt)

        else:
            # Apply the style to it by default
            try:
                run.style = style_name
            except ValueError:
                # Style is not a character style (might be a paragraph style)
                # Skip applying it to the run
                pass
            # para, style_name = handle_code_block_tags_with_suffixes(para, style_name)

            # Apply bolding and/or italicization as necessary
            handled_style_name = style_name
            if handled_style_name.replace(" ", "_") in HANDLE_UNDERSCORE_CLASSES:
                handled_style_name = style_name.replace(" ", "_")

            run.bold = handled_style_name in CSS_BOLD
            run.italic = handled_style_name in CSS_ITALIC

        return run

    tag_regex = re.compile(r"\{{3}(/?[^\{\}]+?)\}{3}")

    def remove_paragraph_content(para: Paragraph):
        """Remove all non-pPr children so content can be rebuilt from parsed tags."""
        for child in list(para._p):
            if child.tag != qn("w:pPr"):
                para._p.remove(child)

    def split_text_into_tagged_segments(text: str):
        """Split text into (segment_text, active_styles) supporting nested tags."""
        segments: list[tuple[str, list[str]]] = []
        active_styles: list[str] = []
        last_pos = 0

        for match in tag_regex.finditer(text):
            if match.start() > last_pos:
                segment_text = text[last_pos : match.start()]
                if segment_text:
                    segments.append((segment_text, active_styles.copy()))

            token = match.group(1)
            if token.startswith("/"):
                closing_name = token[1:]
                # Remove the nearest matching opener from the stack end.
                for idx in range(len(active_styles) - 1, -1, -1):
                    if active_styles[idx] == closing_name:
                        del active_styles[idx]
                        break
            else:
                active_styles.append(token)

            last_pos = match.end()

        if last_pos < len(text):
            trailing = text[last_pos:]
            if trailing:
                segments.append((trailing, active_styles.copy()))

        return segments

    def apply_style_stack(paragraph: Paragraph, text: str, active_styles: list[str]):
        """Create one run/hyperlink and apply all active styles (outer + inline)."""
        if not text:
            return

        link_style = next(
            (style for style in reversed(active_styles) if style.startswith("LINK|")),
            None,
        )

        if link_style:
            url = link_style.split("|", 1)[1]
            run = add_hyperlink(paragraph, url, text)
        else:
            run = paragraph.add_run(text)

        for style_name in active_styles:
            if style_name.startswith("LINK|"):
                continue
            apply_formatting_and_styling(run, style_name)

    def serialize_paragraph_text_preserving_hyperlinks(paragraph: Paragraph):
        """Serialize paragraph text and wrap existing XML hyperlinks as LINK tags."""

        def extract_text_with_controls(node: BaseOxmlElement):
            """Extract text while preserving explicit tabs and line breaks."""
            pieces: list[str] = []
            for part in node.xpath(".//*[self::w:t or self::w:tab or self::w:br]"):
                if part.tag == qn("w:t"):
                    pieces.append(part.text or "")
                elif part.tag == qn("w:tab"):
                    pieces.append("\t")
                elif part.tag == qn("w:br"):
                    pieces.append("\n")
            return "".join(pieces)

        def resolve_hyperlink_target_from_rel_id(rel_id: str | None):
            if not rel_id:
                return None

            rel = None
            try:
                rel = paragraph.part.rels[rel_id]
            except Exception:
                rel = None

            if rel is None:
                return None

            try:
                if rel.reltype != RT.HYPERLINK:
                    return None
            except Exception:
                return None

            return getattr(rel, "target_ref", None)

        def extract_url_from_fldsimple(instr: str):
            if not instr:
                return None

            quoted = re.search(r'HYPERLINK\\s+"([^"]+)"', instr, flags=re.IGNORECASE)
            if quoted:
                return quoted.group(1)

            unquoted = re.search(r"HYPERLINK\\s+([^\\s]+)", instr, flags=re.IGNORECASE)
            if unquoted:
                return unquoted.group(1)

            return None

        chunks: list[str] = []

        for child in paragraph._p:
            if child.tag == qn("w:hyperlink"):
                link_text = extract_text_with_controls(child)
                anchor = child.get(qn("w:anchor"))
                rel_id = child.get(qn("r:id"))

                link_target = None
                if anchor:
                    link_target = f"#{anchor}"
                elif rel_id:
                    link_target = resolve_hyperlink_target_from_rel_id(rel_id)

                if link_target and link_text:
                    chunks.append(
                        f"{{{{{{LINK|{link_target}}}}}}}{link_text}{{{{{{/LINK|{link_target}}}}}}}"
                    )
                else:
                    chunks.append(link_text)

            elif child.tag == qn("w:fldSimple"):
                field_text = extract_text_with_controls(child)
                instr = child.get(qn("w:instr")) or ""
                link_target = extract_url_from_fldsimple(instr)

                if link_target and field_text:
                    chunks.append(
                        f"{{{{{{LINK|{link_target}}}}}}}{field_text}{{{{{{/LINK|{link_target}}}}}}}"
                    )
                else:
                    chunks.append(field_text)

            elif child.tag == qn("w:r"):
                run_text = extract_text_with_controls(child)
                chunks.append(run_text)

        return "".join(chunks)

    # Remove sourceCode tags from EX/NOTE paragraphs (sourceCode is a paragraph style, can't be applied inside EX)
    for paragraph in iter_paragraphs(doc):
        if paragraph.style.name in EXAMPLE_NOTE_CLASSES:
            for run in paragraph.runs:
                run.text = run.text.replace("{{{sourceCode}}}", "").replace(
                    "{{{/sourceCode}}}", ""
                )

    for paragraph in iter_paragraphs(doc):
        full_text = serialize_paragraph_text_preserving_hyperlinks(paragraph)

        if "{{{" not in full_text:
            continue

        segments = split_text_into_tagged_segments(full_text)

        # No tags found in this paragraph: do not alter runs.
        if len(segments) == 1 and segments[0][0] == full_text and not segments[0][1]:
            continue

        remove_paragraph_content(paragraph)

        for segment_text, active_styles in segments:
            apply_style_stack(paragraph, segment_text, active_styles)

    return doc


def normalize_list_item_styles(doc: Doc):
    """
    Normalize styles and indentation for list items within the same list level to ensure
    consistent formatting.

    Pandoc sometimes assigns different styles (Compact vs Normal) to list items in the same
    list, and may also set different explicit w:ind w:left values on individual items.
    This function groups paragraphs by (numId, ilvl) and ensures every item in a group
    shares the same style and the same w:ind w:left value.
    """
    # Group paragraphs by (numId, ilvl)
    lists = {}  # (numId, ilvl) -> [paragraphs]

    for paragraph in iter_paragraphs(doc):
        pPr = paragraph._element.find(qn("w:pPr"))
        if pPr is None:
            continue

        numPr = pPr.find(qn("w:numPr"))
        if numPr is None:
            continue

        numId_elem = numPr.find(qn("w:numId"))
        if numId_elem is None:
            continue

        ilvl_elem = numPr.find(qn("w:ilvl"))
        numId = numId_elem.get(qn("w:val"))
        ilvl = ilvl_elem.get(qn("w:val")) if ilvl_elem is not None else "0"

        key = (numId, ilvl)
        if key not in lists:
            lists[key] = []
        lists[key].append(paragraph)

    # Normalize each (numId, ilvl) group
    for key, paragraphs in lists.items():
        if len(paragraphs) <= 1:
            continue

        # --- Style normalization ---
        styles = {}
        for para in paragraphs:
            style_name = para.style.name
            styles[style_name] = styles.get(style_name, 0) + 1

        target_style = "Compact" if "Compact" in styles else max(styles, key=styles.get)

        for para in paragraphs:
            if para.style.name != target_style:
                try:
                    para.style = target_style
                except Exception:
                    pass

        # --- w:ind w:left normalization ---
        # Collect all explicit w:left values set on individual items in this group.
        left_counts = {}
        for para in paragraphs:
            pPr = para._element.find(qn("w:pPr"))
            if pPr is None:
                continue
            ind = pPr.find(qn("w:ind"))
            if ind is not None:
                left = ind.get(qn("w:left"))
                if left is not None:
                    left_counts[left] = left_counts.get(left, 0) + 1

        if not left_counts:
            continue  # No explicit indentation anywhere in this group — nothing to do.

        # Use the most common explicit left value as the target for the whole group.
        target_left = max(left_counts, key=left_counts.get)

        for para in paragraphs:
            pPr = para._element.find(qn("w:pPr"))
            if pPr is None:
                pPr = OxmlElement("w:pPr")
                para._element.insert(0, pPr)

            ind = pPr.find(qn("w:ind"))
            if ind is None:
                ind = OxmlElement("w:ind")
                pPr.append(ind)

            ind.set(qn("w:left"), target_left)

    return doc


def normalize_continuation_list_paragraphs(doc: Doc):
    """
    Fix indentation AND spacing of Pandoc-generated continuation paragraphs in loose list items.

    When a list item has multiple paragraphs ("loose" list item), Pandoc emits:
    - The first paragraph with the main list numId and an explicit w:ind w:left (and often
      an explicit w:spacing w:after to separate it visually from the next list item).
    - Each continuation paragraph with a *different* continuation numId, no explicit w:ind,
      and no explicit w:spacing.

    This leads to two problems:
    1. Continuation paragraphs are indented differently from the first paragraph of the item.
    2. The space *before* the continuation paragraph (= space after the preceding paragraph)
       is wrong: it uses the "last-item" spacing value rather than the normal intra-item value.

    Fix (two-phase):
    Phase 1  – Collect groups: each group is a primary paragraph followed by its
               continuation paragraphs (consecutive list paragraphs with a different numId
               but the same ilvl, without their own explicit w:left).
    Phase 2  – For each group:
               • Apply the primary paragraph's style to all continuation paragraphs.
               • Copy the primary paragraph's explicit w:ind w:left to all continuations.
               • Move the primary paragraph's explicit w:spacing to the *last* continuation
                 (so Word renders the gap at the end of the whole item, not mid-item).
    """
    paragraphs = list(iter_paragraphs(doc))

    # ---- Phase 1: collect groups ------------------------------------------------
    # Each entry: (primary_para, [continuation_paras], primary_left, primary_numId, ilvl)
    groups: list[tuple] = []

    # Maps ilvl -> (primary_numId, primary_left, primary_para)
    primary_context: dict[str, tuple] = {}

    i = 0
    while i < len(paragraphs):
        para = paragraphs[i]
        pPr = para._element.find(qn("w:pPr"))

        if pPr is None:
            primary_context.clear()
            i += 1
            continue

        numPr = pPr.find(qn("w:numPr"))
        if numPr is None:
            primary_context.clear()
            i += 1
            continue

        numId_elem = numPr.find(qn("w:numId"))
        ilvl_elem = numPr.find(qn("w:ilvl"))
        if numId_elem is None:
            primary_context.clear()
            i += 1
            continue

        numId = numId_elem.get(qn("w:val"))
        ilvl = ilvl_elem.get(qn("w:val")) if ilvl_elem is not None else "0"
        ind = pPr.find(qn("w:ind"))
        explicit_left = ind.get(qn("w:left")) if ind is not None else None

        # Check whether this paragraph is a continuation of the current primary
        if ilvl in primary_context:
            ctx_numId, ctx_left, ctx_primary = primary_context[ilvl]
            if numId != ctx_numId and explicit_left is None:
                # This is a continuation paragraph — collect it under the current primary,
                # then look-ahead to grab any further consecutive continuations.
                continuations = [para]
                j = i + 1
                while j < len(paragraphs):
                    next_para = paragraphs[j]
                    next_pPr = next_para._element.find(qn("w:pPr"))
                    if next_pPr is None:
                        break
                    next_numPr = next_pPr.find(qn("w:numPr"))
                    if next_numPr is None:
                        break
                    next_numId_elem = next_numPr.find(qn("w:numId"))
                    next_ilvl_elem = next_numPr.find(qn("w:ilvl"))
                    if next_numId_elem is None:
                        break
                    next_numId = next_numId_elem.get(qn("w:val"))
                    next_ilvl = (
                        next_ilvl_elem.get(qn("w:val"))
                        if next_ilvl_elem is not None
                        else "0"
                    )
                    next_ind = next_pPr.find(qn("w:ind"))
                    next_left = (
                        next_ind.get(qn("w:left")) if next_ind is not None else None
                    )

                    if (
                        next_numId != ctx_numId
                        and next_ilvl == ilvl
                        and next_left is None
                    ):
                        continuations.append(next_para)
                        j += 1
                    else:
                        break

                groups.append((ctx_primary, continuations, ctx_left))
                # Jump past all collected continuations
                i = j
                continue

        # This is a primary paragraph (has its own explicit w:left, or a fresh item)
        if explicit_left is not None:
            primary_context[ilvl] = (numId, explicit_left, para)
        else:
            primary_context.pop(ilvl, None)

        i += 1

    # ---- Phase 2: apply fixes ---------------------------------------------------
    for primary_para, continuation_paras, target_left in groups:
        primary_pPr = primary_para._element.find(qn("w:pPr"))
        primary_style = primary_para.style.name

        # Lift explicit w:spacing off the primary paragraph and hand it to the last
        # continuation, so visually "item spacing" ends after the whole item, not mid-item.
        spacing_to_transfer = None
        if primary_pPr is not None:
            primary_spacing = primary_pPr.find(qn("w:spacing"))
            if primary_spacing is not None:
                spacing_to_transfer = copy.deepcopy(primary_spacing)
                primary_pPr.remove(primary_spacing)

        for cont_para in continuation_paras:
            cont_pPr = cont_para._element.find(qn("w:pPr"))
            if cont_pPr is None:
                cont_pPr = OxmlElement("w:pPr")
                cont_para._element.insert(0, cont_pPr)

            # Apply the same style as the primary paragraph
            try:
                cont_para.style = primary_style
            except Exception:
                pass

            # Set the same explicit w:left as the primary paragraph
            cont_ind = cont_pPr.find(qn("w:ind"))
            if cont_ind is None:
                cont_ind = OxmlElement("w:ind")
                cont_pPr.append(cont_ind)
            cont_ind.set(qn("w:left"), target_left)

        # Move spacing to the last continuation paragraph
        if spacing_to_transfer is not None:
            last_cont_pPr = continuation_paras[-1]._element.find(qn("w:pPr"))
            if last_cont_pPr is None:
                last_cont_pPr = OxmlElement("w:pPr")
                continuation_paras[-1]._element.insert(0, last_cont_pPr)
            existing = last_cont_pPr.find(qn("w:spacing"))
            if existing is not None:
                last_cont_pPr.remove(existing)
            last_cont_pPr.append(spacing_to_transfer)

    return doc


def set_keep_with_next_false(doc: Doc):
    """By default, `keep with next` is set to `True`. This has the effect of adding blocks of whitespace throughout the document. For all paragraphs, ensure `keep with next` is set to `False` to maintain visual continuity within the document."""

    previous_paragraph = None
    for paragraph in iter_paragraphs(doc):
        
        paragraph.paragraph_format.keep_with_next = False
        
        if previous_paragraph and previous_paragraph.style.name == "Compact" and paragraph.style.name != "Compact":
            
            # 2. Check if the previous paragraph is inside a table
            is_in_table = bool(previous_paragraph._element.xpath('./ancestor::w:tbl'))
            
            # 3. Add 12pt spacing ONLY if the paragraph is NOT in a table.
            # This prevents spacing alterations in the last cell
            if not is_in_table:
                previous_paragraph.paragraph_format.space_after = Pt(12)
                
        previous_paragraph = paragraph

    return doc

def make_vanish_run(marker: str) -> BaseOxmlElement:
        """Return a <w:r> element whose text is *marker* and whose rPr contains <w:vanish/>."""
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rPr.append(OxmlElement("w:vanish"))
        r.append(rPr)
        t = OxmlElement("w:t")
        t.text = marker
        r.append(t)
        return r

def add_sourcecode_vanish_markers(doc: Doc):
    """
    For each group of consecutive 'Source Code' / 'EX Source Code' paragraphs,
    prepend an invisible (w:vanish) run with text 'jsonStartBlockquote' to the first paragraph
    and append an invisible run with text 'jsonEndBlockquote' to the last paragraph.

    This lets callers who read the raw DOCX XML identify code-block runs unambiguously
    and reconstruct the original markdown fenced code blocks on the round-trip back.
    """
    SOURCE_CODE_STYLES = {"Source Code", "EX Source Code"}

    def insert_run_after_pPr(p_elem: BaseOxmlElement, run_elem: BaseOxmlElement):
        """Insert *run_elem* into *p_elem* immediately after <w:pPr>, or at index 0 if absent."""
        pPr = p_elem.find(qn("w:pPr"))
        if pPr is not None:
            idx = list(p_elem).index(pPr)
            p_elem.insert(idx + 1, run_elem)
        else:
            p_elem.insert(0, run_elem)

    paragraphs = list(iter_paragraphs(doc))
    i = 0
    while i < len(paragraphs):
        if paragraphs[i].style.name not in SOURCE_CODE_STYLES:
            i += 1
            continue

        group_start = i
        while i < len(paragraphs) and paragraphs[i].style.name in SOURCE_CODE_STYLES:
            i += 1
        group_end = i - 1

        blockquote_type = None
        first_run = paragraphs[group_start].runs[0] if paragraphs[group_start].runs else None
        last_run = paragraphs[group_end].runs[-1] if paragraphs[group_end].runs else None
        # if first run of the group contains jsonStartBlockquote, it is a json blockquote, if it contains textStartBlockquote, it is a normal text blockquote
        if first_run and "jsonStartBlockquote" in first_run.text:
            blockquote_type = "json"
            first_run.text = first_run.text.replace("jsonStartBlockquote", "")
        elif first_run and "textStartBlockquote" in first_run.text:
            blockquote_type = "text"
            first_run.text = first_run.text.replace("textStartBlockquote", "")
        if last_run and "jsonEndBlockquote" in last_run.text:
            blockquote_type = "json"
            last_run.text = last_run.text.replace("jsonEndBlockquote", "")
        elif last_run and "textEndBlockquote" in last_run.text:
            blockquote_type = "text"
            last_run.text = last_run.text.replace("textEndBlockquote", "")

        insert_run_after_pPr(
            paragraphs[group_start]._p, make_vanish_run(f"{blockquote_type}StartBlockquote" if blockquote_type else "jsonStartBlockquote")
        )
        paragraphs[group_end]._p.append(make_vanish_run(f"{blockquote_type}EndBlockquote" if blockquote_type else "jsonEndBlockquote"))

    return doc


def _run_is_hidden_for_export(run: Run) -> bool:
    """Return True when a run is hidden via w:vanish or InvisibleText style."""
    if run.style and run.style.name == "InvisibleText":
        return True

    run_props = run._element.find(qn("w:rPr"))
    if run_props is None:
        return False

    return run_props.find(qn("w:vanish")) is not None


def build_clean_export_docx(source_docx_path: str | Path) -> Path:
    """Create a temporary DOCX copy without invisible runs, for PDF/EPUB exports."""
    source_path = Path(source_docx_path)
    export_doc = Document(str(source_path))

    for paragraph in iter_paragraphs(export_doc):
        for run in list(paragraph.runs):
            if _run_is_hidden_for_export(run):
                run._element.getparent().remove(run._element)

    cleaned_path = source_path.parent / f"{source_path.stem}.export-clean.docx"
    export_doc.save(str(cleaned_path))
    return cleaned_path


def docx_to_pdf_libreoffice(docx_path, out_dir=None):
    docx_path = Path(docx_path)
    out_dir = Path(out_dir or docx_path.parent)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice 'soffice' not found in PATH")

    # In containers, LibreOffice may fail if HOME/XDG paths are not writable.
    lo_profile_dir = Path(tempfile.mkdtemp(prefix="libreoffice-profile-"))
    lo_cache_dir = Path(tempfile.mkdtemp(prefix="libreoffice-cache-"))
    lo_config_dir = Path(tempfile.mkdtemp(prefix="libreoffice-config-"))

    env = os.environ.copy()
    env["HOME"] = str(lo_profile_dir)
    env["XDG_CACHE_HOME"] = str(lo_cache_dir)
    env["XDG_CONFIG_HOME"] = str(lo_config_dir)

    subprocess.run(
        [
            soffice,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nofirststartwizard",
            "--nolockcheck",
            f"-env:UserInstallation={lo_profile_dir.resolve().as_uri()}",
            "--convert-to",
            "pdf",
            str(docx_path),
            "--outdir",
            str(out_dir),
        ],
        check=True,
        env=env,
    )

    source_pdf = out_dir / f"{docx_path.stem}.pdf"
    target_pdf = out_dir / f"{docx_path.stem.removesuffix('.export-clean')}.pdf"
    if source_pdf.exists() and source_pdf != target_pdf:
        source_pdf.replace(target_pdf)


def docx_to_epub(docx_path, out_dir=None):
    docx_path = Path(docx_path)
    out_dir = Path(out_dir or docx_path.parent)
    target_epub = out_dir / f"{docx_path.stem.removesuffix('.export-clean')}.epub"

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice 'soffice' not found in PATH")

    lo_profile_dir = Path(tempfile.mkdtemp(prefix="libreoffice-profile-"))
    lo_cache_dir = Path(tempfile.mkdtemp(prefix="libreoffice-cache-"))
    lo_config_dir = Path(tempfile.mkdtemp(prefix="libreoffice-config-"))

    env = os.environ.copy()
    env["HOME"] = str(lo_profile_dir)
    env["XDG_CACHE_HOME"] = str(lo_cache_dir)
    env["XDG_CONFIG_HOME"] = str(lo_config_dir)

    subprocess.run(
        [
            soffice,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nofirststartwizard",
            "--nolockcheck",
            f"-env:UserInstallation={lo_profile_dir.resolve().as_uri()}",
            "--convert-to",
            "epub",
            str(docx_path),
            "--outdir",
            str(out_dir),
        ],
        check=True,
        env=env,
    )

    source_epub = out_dir / f"{docx_path.stem}.epub"
    if source_epub.exists() and source_epub != target_epub:
        source_epub.replace(target_epub)


def parse_requested_outputs(extra_outputs: str | None) -> set[str]:
    """Parse extra output selection, defaulting to PDF for backward compatibility."""
    if extra_outputs is None:
        return {"pdf"}
    if not extra_outputs.strip():
        return set()
    return {part.strip() for part in extra_outputs.split(",") if part.strip()}


# endregion


def postprocess(docx_dir: str, extra_outputs: str | None = None):
    """
    ### Description
    Performs postprocessing on the generated docx:
    1. Formats the references to ensure the reference body is indented with respect to its label.
    2. Formats tables to ensure their border widths are of 1/2 pt thickness
    3. Adds indentation between example/note tags and their bodies
    4. Formats various tagged portions of text with the appropriate formatting, since Pandoc doesn't handle these sections of text well.
    5. Removes excess spaces after paragraphs

    ### Arguments
    - `docx_dir`: The absolute or relative path at which the generated Docx was saved
    """
    doc: Doc = Document(docx_dir)

    doc = remove_startEXNO_invisible_text_when_in_the_same_paragraph(doc)
    doc = format_references(doc)
    doc = format_tables(doc)
    doc = format_examples_and_notes(doc)
    doc = replace_whitespace_placeholders_with_chars(doc)
    doc = add_tagged_styles_and_formatting(doc)

    doc = normalize_list_item_styles(doc)
    doc = normalize_continuation_list_paragraphs(doc)
    doc = set_keep_with_next_false(doc)
    doc = add_sourcecode_vanish_markers(doc)
    doc = merge_example_with_next_sourcecode(doc)

    doc.save(docx_dir)


    containing_folder = os.path.dirname(docx_dir)
    include_frontpage(containing_folder)

    requested_outputs = parse_requested_outputs(extra_outputs)

    # clean up any old generated PDF/EPUB files before generating new ones
    pdf_path = Path(docx_dir).parent / "document.pdf"
    if pdf_path.exists():
        pdf_path.unlink()

    epub_path = Path(docx_dir).parent / "document.epub"
    if epub_path.exists():
        epub_path.unlink()

    if requested_outputs:
        export_docx_path = build_clean_export_docx(docx_dir)
        try:
            if "pdf" in requested_outputs:
                docx_to_pdf_libreoffice(export_docx_path)
            if "epub" in requested_outputs:
                docx_to_epub(export_docx_path)
        finally:
            if export_docx_path.exists():
                export_docx_path.unlink()